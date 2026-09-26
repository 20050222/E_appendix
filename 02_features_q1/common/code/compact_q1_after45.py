from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from enrich_q1_docx import configure_styles, set_cell_text, style_table


INPUT = Path(r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_25\4问题一求解及分析_正文增强稿.docx")
OUTPUT = INPUT.with_name("4问题一求解及分析_正文精简稿.docx")


def set_table_pagination(table) -> None:
    if not table.rows:
        return
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    header_trpr.append(header)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def append_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(8)


def append_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], source: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    r = p.add_run(caption)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)

    table = doc.add_table(rows=1, cols=len(headers))
    for i, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF", size=8.8)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.8)
    style_table(table)
    set_table_pagination(table)
    append_source_note(doc, source)


def remove_body_after_heading(doc: Document, needle: str) -> None:
    body = doc.element.body
    target = None
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            text = "".join(child.itertext())
            if needle in text:
                target = child
                break
    if target is None:
        raise ValueError(f"找不到章节锚点: {needle}")
    remove = False
    for child in list(body.iterchildren()):
        if child is target:
            remove = True
        if remove and child.tag != qn("w:sectPr"):
            body.remove(child)


def add_body(doc: Document, text: str, lead: str | None = None) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if lead and text.startswith(lead):
        r = p.add_run(lead)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        text = text[len(lead):]
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)


def main() -> None:
    doc = Document(str(INPUT))
    configure_styles(doc)
    remove_body_after_heading(doc, "4.6 结果与质量核验")

    doc.add_heading("4.6 结果复核", level=2)
    add_body(doc, "本节只保留与题目验收条件直接相关的三组证据：样本是否完整、特征结构是否合法、对齐接口和标签是否一致。覆盖率用于描述可观测性，不作为情感分类或回归性能指标。", lead="复核口径。")
    add_body(doc, "100 条样本均建立 sample_id 和 aligned 文件；文本、音频均保留 100 条，视觉有 90 条样本至少包含一个有效人脸桶，另外 10 条明确记录为 no_valid_face_frames。失败样本没有从索引中删除，而是以零填充、vision_mask=0 和状态字段保留，因此文件完整性与模态可观测性被分开统计。", lead="样本完整性。")

    append_table(doc, "表 4-12 三模态特征覆盖与结构核验", ["模态", "输出形状", "有效桶/总桶", "覆盖率", "至少一个有效桶", "结构核验"], [
        ["文本", "50×768", "2259/5000", "45.18%（序列位置占用率）", "100/100", "100/100 PASS"],
        ["音频", "50×74", "4957/5000", "99.14%（视频时间桶）", "100/100", "100/100 PASS"],
        ["视觉", "50×35", "3865/5000", "77.30%（视频时间桶）", "90/100", "100/100 接口通过"],
    ], "数据来源：stepE_alignment/data/multimodal_coverage_summary.csv、stepF_validation/data/final_sample_validation.csv。文本位置不是视频时间轴，三种覆盖率不能直接解释为识别准确率。")

    add_body(doc, "数值检查显示，三种模态均为 float32，有限值比例均为 100%，mask 非法样本数均为 0，且不存在 mask=1 但特征整行全零的情况。文本的全零行来自短文本分桶后的空位置；视觉的全零行主要来自无脸桶和 10 条视觉失败样本，均可由对应 mask 与状态字段解释。", lead="数值合法性。")

    append_table(doc, "表 4-13 官方接口与标签交叉核验", ["核验项目", "结果"], [
        ["官方样本数；train/valid/test", "4850；3395 / 728 / 727"],
        ["官方接口形状", "text N×50×768；audio N×50×74；vision N×50×35"],
        ["共享 sample_id", "18"],
        ["分类标签不一致数；回归标签不一致数", "0；0"],
    ], "数据来源：stepF_validation/data/official_interface_comparison.json。该核验比较字段、形状、ID 和标签口径，不声称逐元素复现官方特征数值。")

    add_body(doc, "以上结果满足问题一的三项要求：sample_id、模态文件和特征输出可以一一回溯；50 个位置的边界、source_count、valid 和 mask 记录了序列组织与填充规则；模型版本、采样参数、环境文件、配置、代码和 CSV 日志构成了可复现材料。需要保留的主要风险是视觉模态存在 10 条整段不可观测样本，以及文本桶属于 token 序列位置而非词级真实时间戳。", lead="结论与风险。")

    doc.add_heading("4.7 本问小结", level=2)
    add_body(doc, "问题一建立了以 sample_id 为主键的三模态特征工程流程：文本由 BERT 生成 50×768 表示，音频由 16 kHz 帧级声学特征生成 50×74 表示，视觉由人脸关键点、边框和灰度信息生成 50×35 表示；音频和视觉按视频相对时间划分 50 个桶，文本按 Token 顺序组织到 50 个位置。每个模态均配套独立 mask、有效长度、源观测数和状态字段。")
    add_body(doc, "因此，本问的交付对象应按‘索引与标签—单模态特征—aligned-50 对齐接口—质量与复现材料’归档。后续问题可直接读取 aligned-50，并在模拟模态缺失时新增 synthetic_missing_mask，不能覆盖原始观测 mask。")

    doc.core_properties.title = "问题一求解及分析"
    doc.core_properties.subject = "复杂场景下多模态情感识别的特征建模与时间组织"
    doc.core_properties.comments = "压缩 4.6 结果复核和 4.7 小结，保留题目验收所需的核心统计与风险。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
