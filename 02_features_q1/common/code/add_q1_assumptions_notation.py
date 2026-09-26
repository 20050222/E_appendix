from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from enrich_q1_docx import configure_styles, set_cell_text, style_table


INPUT = Path(r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_26\论文-第一问数据处理.docx")
OUTPUT = INPUT.with_name("论文-第一问数据处理_模型假设与符号说明稿.docx")


def set_table_pagination(table) -> None:
    if not table.rows:
        return
    trpr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)
    for row in table.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))


def set_paragraph_text(paragraph, text: str, *, bold: bool = False, size: float = 10.5) -> None:
    for run in paragraph.runs:
        run.text = ""
    if paragraph.runs:
        run = paragraph.runs[0]
    else:
        run = paragraph.add_run()
    run.text = text
    run.bold = bold
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)


def insert_body_before(anchor, text: str, lead: str | None = None):
    p = anchor.insert_paragraph_before()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if lead and text.startswith(lead):
        r = p.add_run(lead)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        text = text[len(lead):]
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)
    return p


def insert_heading_before(anchor, text: str, level: int = 2):
    p = anchor.insert_paragraph_before(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def insert_table_before(anchor, caption: str, headers: list[str], rows: list[list[str]], source: str | None = None):
    cap = anchor.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)

    table = anchor._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    anchor._p.addprevious(table._element)
    for i, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF", size=9)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=9)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    style_table(table)
    set_table_pagination(table)
    if source:
        note = anchor.insert_paragraph_before(source)
        note.paragraph_format.space_after = Pt(6)
        for run in note.runs:
            run.italic = True
            run.font.name = "Microsoft YaHei"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
            run.font.size = Pt(8)
    return table


def main() -> None:
    doc = Document(str(INPUT))
    configure_styles(doc)

    # Locate the existing blank chapter-2 placeholders.
    chapter = next(p for p in doc.paragraphs if p.text.strip() == "模型假设与符号说明")
    q1_assumptions = next(p for p in doc.paragraphs if p.text.strip() == "模型假设" and p.style.name.startswith("Heading"))
    notation = next(p for p in doc.paragraphs if p.text.strip() == "定义与符号说明")
    dataset = next(p for p in doc.paragraphs if p.text.strip() == "数据集介绍与可视化")

    set_paragraph_text(chapter, "2 模型假设与符号说明", bold=True, size=16)
    set_paragraph_text(q1_assumptions, "2.1 第一问模型假设", bold=True, size=12)
    set_paragraph_text(notation, "2.2 第一问符号说明", bold=True, size=12)

    # Remove generic placeholder paragraphs between the existing headings.
    for p in list(doc.paragraphs):
        if p is q1_assumptions or p is notation or p is dataset:
            continue
        if p.text.strip() in {
            "例如：时间桶边界按视频时长确定；缺失掩码能正确反映特征不可用状态。每条假设都应说明依据及可能影响。",
            "定义样本、模态、时间位置、特征、缺失掩码、情感极性、情感强度和模态贡献等符号。",
        }:
            element = p._element
            element.getparent().remove(element)

    assumptions = [
        ("假设 1：样本主键唯一且可回溯。", "每条视频样本由 video_id 与 clip_id 唯一确定，统一写为 sample_id。后续文本、音频、视觉特征和标签均通过 sample_id 连接，不依赖文件排序或数组下标。若主键重复或视频路径存在歧义，该样本应保留在索引中并标记失败原因。"),
        ("假设 2：附件 1 转写可作为文本输入。", "题目提供的 text 字段视为当前版本的文本观测，不重新运行 ASR 替换它。预处理只做 Unicode 规范化、首尾空白清除和连续空白折叠，否定词、数字、缩写及有语义标点不被静默删除。该假设减少了 ASR 模型和解码参数引入的额外不确定性。"),
        ("假设 3：文本和视频时间采用不同的位置语义。", "文本缺少可靠的词级时间戳，因此将 Token 按顺序均匀分配到 50 个序列位置；音频和视觉则按照真实帧时间映射到视频相对时间桶。三种模态共享序列长度，但不把文本第 k 位解释为视频中的精确发生时刻。"),
        ("假设 4：视频时间可以划分为 50 个归一化区间。", "对第 i 条视频，以其读取到的有效时长 D_i 为边界，将区间划分为 50 个左闭右开桶，最后一个桶覆盖到 D_i。该处理使不同时长、不同原始帧率的视频具有统一的数组接口，同时保留每个桶的起止时间。"),
        ("假设 5：桶内特征用有效观测的均值表示。", "落入同一桶的有效音频帧或有效人脸帧使用算术均值聚合；没有有效观测的桶填充零向量并将对应 mask 置为 0。零向量仅是存储填充值，不被视为真实的静音、无表情或中性情绪。"),
        ("假设 6：视觉主体可由最大有效人脸框确定。", "每个采样帧独立检测人脸，在允许的检测范围内选择 landmark 外接框面积最大的人脸作为主体；检测不到有效人脸时不沿用上一帧特征，而是记录失败并令 vision_mask=0。该规则保证同一配置重跑时主体选择确定。"),
        ("假设 7：标签、特征和状态字段分别承担不同职责。", "annotation 表示离散情感类别，label 表示连续情感强度；特征张量表示数值输入，mask 和 status 表示观测是否可用。覆盖率、有效桶数和结构通过率只用于数据质量核验，不直接等同于分类准确率或回归误差。"),
        ("假设 8：原始数据只读且处理结果可独立复现。", "00_raw_readonly 中的原始视频、标签和附件不被覆盖；清洗、截断、掩码和标准化结果写入后续目录。模型版本、关键参数、随机种子、运行环境、代码和逐样本日志共同构成复现条件。"),
    ]

    insert_body_before(notation, "第一问的假设只约束特征提取、时序组织和数据验收，不预设问题二、问题三的预测模型形式。假设的作用是把原始媒体转化为可计算、可回溯的统一输入，并明确哪些结论来自数据工程、哪些结论需要后续预测实验验证。", lead="说明。")
    for lead, body in assumptions:
        insert_body_before(notation, lead + body, lead=lead)

    symbols = [
        ["样本与数据集", "i, N", "样本索引；N 为当前数据集样本数，第一问附件 1 中 N=100。"],
        ["主键", "s_i", "第 i 条样本的 sample_id，由 video_id 与 clip_id 规范化拼接得到。"],
        ["模态", "m∈{T,A,V}", "T、A、V 分别表示文本、音频和视觉模态。"],
        ["统一位置", "K, k", "K 为公共位置数，本文 K=50；k=0,1,…,K−1 为位置或时间桶编号。"],
        ["视频时长", "D_i", "第 i 条视频读取到的有效时长，单位为秒。"],
        ["视频时间桶", "I_{ik}", "第 i 条视频的第 k 个归一化时间区间。"],
        ["文本 Token", "M_i, j, b_i(j)", "M_i 为有效 Token 数；j 为 Token 编号；b_i(j) 为 Token 所属序列桶。"],
        ["文本表示", "h_{ij}, X_{ik}^{(T)}", "h_{ij} 为 BERT Token 表示；X_{ik}^{(T)} 为第 k 个文本桶的 768 维表示。"],
        ["音频帧特征", "a_{ir}, X_{ik}^{(A)}", "a_{ir} 为第 r 个音频帧的 74 维特征；X_{ik}^{(A)} 为音频桶表示。"],
        ["视觉帧特征", "v_{iq}, X_{ik}^{(V)}", "v_{iq} 为第 q 个有效人脸帧的 35 维特征；X_{ik}^{(V)} 为视觉桶表示。"],
        ["源观测集合", "𝒥_{ik}, 𝒵_{ik}^{(m)}", "分别表示文本 Token 集合和模态 m 在时间桶内的有效源观测集合。"],
        ["观测数量", "n_{ik}^{(m)}", "第 i 条样本、第 m 个模态、第 k 个桶内的有效源观测数量。"],
        ["观测掩码", "M_{ik}^{(m)}", "二值观测状态；1 表示存在有效观测，0 表示空桶、无效帧或处理失败。"],
        ["标签", "y_i^{(c)}, y_i^{(r)}", "离散情感类别与连续情感强度，分别对应 annotation 与 label。"],
        ["时间与状态", "t, status, failure_reason", "t 为帧或音频观测时间；status 为阶段状态；failure_reason 为失败原因。"],
    ]
    insert_body_before(notation, "符号说明仅覆盖第一问的特征提取、50 桶组织和质量核验。问题二和问题三可能引入新的缺失机制、预测函数、损失函数、解释指标和测试集符号，暂不与第一问共用未定义符号。", lead="使用范围。")
    insert_table_before(notation, "表 2-1 第一问主要符号说明", ["类别", "符号", "含义"], symbols, "符号与数据字段对应关系：stepA_index、stepB_text、stepC_audio、stepD_vision、stepE_alignment 的配置、schema 和日志文件。")

    # Reserve model-assumption and notation blocks for the two unfinished questions.
    insert_heading_before(dataset, "2.3 第二问模型假设（预留）", level=2)
    insert_body_before(dataset, "本节待补充：局部模态缺失的生成机制、缺失位置与缺失时长的定义、训练与验证数据边界、鲁棒预测模型可接受的观测条件，以及对附件 3 推理样本的适用范围。具体假设应以队友完成的问题二模型和实验设计为准。", lead="待补充内容。")
    insert_heading_before(dataset, "2.4 第二问符号说明（预留）", level=2)
    insert_table_before(dataset, "表 2-2 第二问符号说明（待补充）", ["类别", "符号", "含义/填写要求"], [
        ["缺失机制", "待定", "补充缺失模态、缺失区间、缺失比例及 synthetic_missing_mask 的符号。"],
        ["预测模型", "待定", "补充分类器、回归器、融合函数、损失函数和超参数。"],
        ["评价指标", "待定", "补充 Accuracy、F1、MAE、相关系数等指标的符号和计算口径。"],
    ], "预留表：待问题二模型确定后替换，不应直接沿用第一问符号。")
    insert_heading_before(dataset, "2.5 第三问模型假设（预留）", level=2)
    insert_body_before(dataset, "本节待补充：可解释性模型的决策对象、解释范围、模态贡献定义、局部证据定位规则、解释稳定性检验和附件 4 推理输出条件。具体内容应以队友完成的问题三模型为准。", lead="待补充内容。")
    insert_heading_before(dataset, "2.6 第三问符号说明（预留）", level=2)
    insert_table_before(dataset, "表 2-3 第三问符号说明（待补充）", ["类别", "符号", "含义/填写要求"], [
        ["解释对象", "待定", "补充预测结果、局部片段、时间桶和解释样本的符号。"],
        ["模态贡献", "待定", "补充模态重要性、消融差异、贡献度或权重的定义。"],
        ["证据定位", "待定", "补充文本、语音和视觉证据回溯到原始素材的符号。"],
    ], "预留表：待问题三模型确定后替换，不提前假设解释算法。")

    # Keep table headers on subsequent pages.
    for table in doc.tables:
        if table.rows:
            set_table_pagination(table)

    doc.core_properties.title = "第一问数据处理与模型假设符号说明"
    doc.core_properties.subject = "问题一模型假设、符号说明及后续问题预留结构"
    doc.core_properties.comments = "在论文-第一问数据处理初稿基础上补充第一问假设和符号，预留问题二、问题三位置。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
