from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


INPUT = Path(r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_26\论文-第一问数据处理_模型假设与符号说明稿.docx")
OUTPUT = INPUT.with_name("论文-第一问数据处理_模型假设与符号说明稿.docx")


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def move_before(element, reference) -> None:
    reference.addprevious(element)


def main() -> None:
    doc = Document(str(INPUT))
    paragraphs = doc.paragraphs
    chapter = next(p for p in paragraphs if p.style.name == "Heading 1" and "模型假设与符号说明" in p.text)
    q1 = next(p for p in paragraphs if p.style.name == "Heading 2" and "第一问模型假设" in p.text)
    notation = next(p for p in paragraphs if p.style.name == "Heading 2" and "第一问符号说明" in p.text)
    first_assumption_body = next(p for p in paragraphs if p.text.startswith("第一问的假设只约束"))
    notation_intro = next(p for p in paragraphs if p.text.startswith("符号说明仅覆盖"))
    q2_assumption = next(p for p in paragraphs if p.style.name == "Heading 2" and "第二问模型假设" in p.text)
    q2_notation = next(p for p in paragraphs if p.style.name == "Heading 2" and "第二问符号说明" in p.text)
    q3_assumption = next(p for p in paragraphs if p.style.name == "Heading 2" and "第三问模型假设" in p.text)
    q3_notation = next(p for p in paragraphs if p.style.name == "Heading 2" and "第三问符号说明" in p.text)

    # Heading styles already supply chapter numbering; keep only the title text.
    set_text(chapter, "模型假设与符号说明")
    set_text(q1, "第一问模型假设")
    set_text(notation, "第一问符号说明")
    # These headings were initially written with manual section numbers while
    # using Word's numbered Heading 2 style. Remove the manual prefixes so
    # Word contributes each number exactly once.
    set_text(q2_assumption, "第二问模型假设（预留）")
    set_text(q2_notation, "第二问符号说明（预留）")
    set_text(q3_assumption, "第三问模型假设（预留）")
    set_text(q3_notation, "第三问符号说明（预留）")

    # Put the chapter heading before 2.1 and 2.2 before the inserted content.
    move_before(chapter._element, q1._element)
    move_before(notation._element, notation_intro._element)

    # Remove the old generic heading that preceded the real chapter-2 heading.
    for p in list(doc.paragraphs):
        if p.text.strip() == "模型假设" and p.style.name.startswith("Heading"):
            remove_paragraph(p)

    # Repair the static table-of-contents cache left by the source document.
    # The inserted headings are listed in the same location as the original
    # chapter-2 placeholders; the duplicate TOC line near the body is removed.
    toc = list(doc.paragraphs)
    toc_15 = next(p for p in toc if p.style.name.lower() == "toc 2" and p.text.startswith("1.5"))
    toc_old_q1 = next(p for p in toc if p.style.name.lower() == "toc 2" and p.text.startswith("2.1"))
    toc_old_notation = next(p for p in toc if p.style.name.lower() == "toc 2" and p.text.startswith("2.2"))
    set_text(toc_15, "1.5 模型假设\t5")
    set_text(toc_old_q1, "2.1 第一问模型假设\t6")
    set_text(toc_old_notation, "2.2 第一问符号说明\t6")
    for p in list(doc.paragraphs):
        if p.style.name.lower() == "toc 1" and "模型假设与符号说明" in p.text:
            remove_paragraph(p)
    toc_old_q1.insert_paragraph_before("2 模型假设与符号说明\t6", style="toc 1")
    chapter3_toc = next(p for p in toc if p.style.name.lower() == "toc 1" and p.text.startswith("3 "))

    # Remove blank paragraphs between the chapter heading and the first substantive paragraph.
    paragraphs = doc.paragraphs
    first_nonblank_after = next(p for p in paragraphs if p.text.startswith("第一问的假设只约束"))
    in_section = False
    for p in list(doc.paragraphs):
        if p._element is chapter._element:
            in_section = True
            continue
        if p._element is first_nonblank_after._element:
            break
        if in_section and not p.text.strip() and not p._p.xpath('.//w:drawing'):
            remove_paragraph(p)

    doc.save(str(OUTPUT))
    print(OUTPUT)


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    paragraph.add_run(text)


if __name__ == "__main__":
    main()
