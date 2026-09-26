from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


WORKSPACE = Path(r"E:\myProject\postgraduate\Coding\CPMCM\E_workspace")
SOURCE = Path(r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_25\4问题一求解及分析_章节草稿.docx")
OUTPUT = SOURCE.with_name("4问题一求解及分析_章节修订稿.docx")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "D9D9D9", size: str = "4") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_text(cell, text: str, *, bold: bool = False, color: str = "000000", size: float = 9.2) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def style_table(table, header_fill: str = "365F8D") -> None:
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, header_fill)
                for run in cell.paragraphs[0].runs:
                    run.font.color.rgb = RGBColor(255, 255, 255)
                    run.bold = True
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F4F7FA")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_caption(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = "Caption" if "Caption" in [s.name for s in doc.styles] else "Normal"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)


def add_source_note(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    r.italic = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(8)
    r.font.color.rgb = RGBColor(100, 100, 100)


def add_figure(doc: Document, image: Path, caption: str, source: str, width: float = 6.15) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(image), width=Inches(width))
    add_caption(doc, caption)
    add_source_note(doc, source)


def add_table(doc: Document, caption: str, headers: list[str], rows: list[list[str]], source: str) -> None:
    add_caption(doc, caption)
    table = doc.add_table(rows=1, cols=len(headers))
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, color="FFFFFF")
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            set_cell_text(cells[j], value)
    style_table(table)
    add_source_note(doc, source)


def replace_text(doc: Document, old: str, new: str) -> int:
    count = 0
    for p in doc.paragraphs:
        if old in p.text:
            for run in p.runs:
                run.text = ""
            p.runs[0].text = new if p.runs else p.add_run(new).text
            count += 1
    return count


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def find_drawing_paragraphs(doc: Document):
    return [p for p in doc.paragraphs if p._p.xpath(".//w:drawing")]


def configure_styles(doc: Document) -> None:
    for style_name in ("Normal", "Body Text"):
        if style_name not in [s.name for s in doc.styles]:
            continue
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)
        style.paragraph_format.line_spacing = 1.35
        style.paragraph_format.space_after = Pt(6)
    for style in doc.styles:
        if style.name.startswith("Heading") or style.name == "Title":
            style.font.color.rgb = RGBColor(0, 0, 0)
            style.font.name = "Microsoft YaHei"
            style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")


def main() -> None:
    doc = Document(str(SOURCE))
    configure_styles(doc)

    replace_text(doc, "问题一求解及分析章节草稿", "问题一求解及分析")
    replace_text(doc, "本章草稿只覆盖问题一。它把队友的工程记录重组为“输入—表示—对齐—输出—核验”的论文逻辑，不包含问题二和问题三的预测模型，也不把特征覆盖率当作分类或回归性能。", "本章围绕问题一建立文本、音频和视觉三模态特征提取与时序组织流程。章节依次说明样本索引、模态特征表示、统一 50 桶接口、质量核验和复现记录，并将每个结论对应到可追溯的数据文件和图表。特征覆盖率用于描述输入质量，不作为分类或回归性能指标。")
    replace_text(doc, "现有 text_index.csv 可直接核验 100 条文本索引，text_feature_validation.csv 中 100 条记录均标记为 PASS。队友运行笔记进一步记录：标签行数、MP4 数和唯一 sample_id 数均为 100，缺失视频和歧义匹配均为 0；这些汇总值在正式提交前仍需与原始 sample_index.csv 复算。", "工作区复核得到：标签行数、MP4 数和唯一 sample_id 数均为 100，缺失视频和歧义匹配均为 0，视频元信息可读率为 100/100；text_feature_validation.csv 中 100 条记录均为 PASS。")
    replace_text(doc, "表 4-5 三模态位置覆盖汇总 统计值待原始结果文件复算", "表 4-5 三模态位置覆盖汇总。完整统计来自 multimodal_coverage_summary.csv。")
    replace_text(doc, "表 4-6 三类边界样本的桶级回溯摘要 统计值待原始结果文件复算", "表 4-6 三类边界样本的桶级回溯摘要。完整记录来自 representative_traceability.csv。")
    replace_text(doc, "图 4-3 三模态位置覆盖与样本级可观测性 统计值待原始结果文件复算", "图 4-3 三模态总体质量概览。图中同时展示样本级有效桶分布、总体覆盖率以及至少一个有效桶和 50 桶全有效样本数。")
    replace_text(doc, "该接口可以直接用于问题二的局部模态缺失建模和问题三的证据定位。", "该接口可以直接作为问题二的输入。若后续构造人为缺失场景，应新增 synthetic_missing_mask，不能覆盖问题一记录的原始模态 mask。")
    replace_text(doc, "且本次收到的材料缺少音频、视觉和最终验收原始文件。", "当前工作区已经保存音频、视觉和最终验收文件；正式提交时仍应将完整特征文件、日志、配置和代码按赛事要求一并归档。")

    if doc.tables:
        set_cell_text(doc.tables[0].cell(0, 0), "核验说明  本修订稿使用当前 E_workspace 中的 sample_index、三模态特征、aligned-50、质量核验和复现文件生成。所有统计数字均可由对应 CSV、JSON 或 NPZ 文件复核。")
        set_cell_text(doc.tables[3].cell(0, 0), "实现口径  视觉主体选择采用实际运行规则：每个采样帧独立选择 landmark 外接框面积最大的人脸。检测不到人脸时不沿用上一帧特征，并将 vision_mask 置为 0。")

    drawings = find_drawing_paragraphs(doc)
    if len(drawings) >= 3:
        remove_paragraph(drawings[2])
        captions = [p for p in doc.paragraphs if "图 4-3 三模态总体质量概览" in p.text]
        if captions:
            anchor = captions[0]
            image = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig1_multimodal_qc_overview/fig1_multimodal_qc_overview.png"
            image_p = anchor.insert_paragraph_before()
            image_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            image_p.add_run().add_picture(str(image), width=Inches(6.15))
            source_p = anchor.insert_paragraph_before("数据来源：stepE_alignment/data/multimodal_coverage_summary.csv；中文出版图由 plot_multimodal_publication_chinese.py 生成。")
            source_p.paragraph_format.space_after = Pt(5)

    # Add verified quality tables and a complete figure appendix after the existing chapter.
    doc.add_page_break()
    h = doc.add_paragraph("4.8 图表证据与电子材料对应关系", style="Heading 2")
    h.paragraph_format.keep_with_next = True
    doc.add_paragraph("正文图表用于解释方法和总体结果，附录图用于展示逐桶 mask、标签分层、流程追踪、最终验收和复现证据。每张图均保留 PNG、PDF 和 SVG 三种格式，图源代码位于对应阶段的 code 目录。提取诊断图仅作为附录/备查材料保留，不在正文或自动附录图清单中插入。")

    add_table(doc, "表 4-7 数值质量核验摘要", ["模态", "dtype", "有效值比例", "mask 非法样本", "全零行", "mask=1 且全零行", "相邻有效桶 L2 变化中位数"], [
        ["文本", "float32", "100%", "0", "2741", "0", "14.0345"],
        ["音频", "float32", "100%", "0", "43", "0", "1274.8173"],
        ["视觉", "float32", "100%", "0", "1135", "0", "0.0812"],
    ], "数据来源：stepF_validation/data/numeric_quality_summary.csv。全零行主要位于 mask=0 的不可观测位置。")

    add_table(doc, "表 4-8 官方接口交叉核验", ["项目", "结果"], [
        ["官方样本总数", "4850"],
        ["train / valid / test", "3395 / 728 / 727"],
        ["text / audio / vision 形状", "N×50×768 / N×50×74 / N×50×35"],
        ["共享 sample_id", "18"],
        ["分类标签不一致数", "0"],
        ["回归标签不一致数", "0"],
    ], "数据来源：stepF_validation/data/official_interface_comparison.json。该核验比较接口、ID 和标签口径，不比较两套特征的逐元素数值。")

    add_table(doc, "表 4-9 复现材料清单", ["材料", "文件", "作用"], [
        ["处理配置", "processing_config.yaml", "固定特征维度、时间桶、mask 和随机种子"],
        ["统一日志", "unified_extraction_log.csv", "记录 A–E 阶段逐样本状态"],
        ["复现清单", "reproducibility_manifest.json", "记录关键文件和代码 SHA-256"],
        ["环境", "environment_cpu.yaml / environment_cuda4090.yaml", "记录 CPU 与 RTX 4090 环境"],
        ["源代码", "stepA–stepG/code/*.py", "固定处理算法与运行顺序"],
    ], "数据来源：stepG_repro/data 和 01_manifest/env。当前清单尚未登记全部单样本 NPZ 的逐文件哈希。")

    doc.add_paragraph("附录图的使用说明", style="Heading 3")
    doc.add_paragraph("附录中的图 4-4 至图 4-13用于补充正文证据。文本图横轴表示 token 序列位置；音频和视觉图横轴表示视频相对时间桶。图中的覆盖率、可用桶数和验收状态均为数据工程质量指标，不代表情感识别准确率。提取诊断图另行保留为附录/备查材料，不参与本章节自动排版。")

    figures = [
        (WORKSPACE / "02_features_q1/stepB_text/image/text_token_count_distribution.png", "图 4-4 文本内容 token 数分布", "数据来源：stepB_text/image/text_token_count_distribution.png。用于说明文本长度分布及中位数 21。"),
        (WORKSPACE / "02_features_q1/stepB_text/image/text_token_count_by_annotation.png", "图 4-5 按情感标签分层的文本 token 数", "数据来源：stepB_text/image/text_token_count_by_annotation.png。该图为描述性统计，不表示标签预测性能。"),
        (WORKSPACE / "02_features_q1/stepB_text/image/text_mask_coverage.png", "图 4-6 文本 50 个 token 序列位置的 mask", "数据来源：stepB_text/image/text_mask_coverage.png。空桶表示 token 序列位置没有分配到内容 token。"),
        (WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig2_mask_structure_revised/fig2_mask_structure_revised.png", "图 4-7 按正确序列语义显示的三模态观测 mask", "数据来源：stepE_alignment/data/multimodal_sample_audit.csv；文本轴不是视频时间轴。"),
        (WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig3_alignment_profiles/fig3_alignment_profiles.png", "图 4-8 三模态逐桶可用率与音视频联合 mask", "数据来源：stepE_alignment/data/multimodal_alignment_log.csv。用于定位逐桶缺失位置。"),
        (WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig4_label_stratified_availability/fig4_label_stratified_availability.png", "图 4-9 按情感标签分层的三模态有效桶数", "数据来源：stepE_alignment/data/multimodal_sample_audit.csv。结果仅用于检查可观测性是否随标签分布而变化。"),
        (WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig6_pipeline_traceability/fig6_pipeline_traceability.png", "图 4-10 样本在 Step A–E 流程中的保留情况", "数据来源：stepG_repro/data/unified_extraction_log.csv。失败样本保留并由状态字段和 mask 标记。"),
        (WORKSPACE / "02_features_q1/stepF_validation/image/publication_chinese/fig7_final_validation_audit/fig7_final_validation_audit.png", "图 4-11 aligned-50 特征集最终验收审计", "数据来源：stepF_validation/data/final_sample_validation.csv、numeric_quality_summary.csv 和 official_interface_comparison.json。"),
        (WORKSPACE / "02_features_q1/stepF_validation/image/representative_traceability.png", "图 4-12 代表性样本的 50 桶回溯", "数据来源：stepF_validation/data/representative_traceability.csv。该图证明位置可回溯，不证明文本桶具有真实词级时间戳。"),
        (WORKSPACE / "02_features_q1/stepG_repro/image/publication_chinese/fig8_reproducibility_record_audit/fig8_reproducibility_record_audit.png", "图 4-13 问题一流程的复现记录覆盖情况", "数据来源：stepG_repro/data/reproducibility_manifest.json、unified_extraction_log.csv 和 processing_config.yaml。")
    ]
    for image, caption, source in figures:
        if image.is_file():
            add_figure(doc, image, caption, source)

    doc.add_paragraph("电子材料交付清单", style="Heading 3")
    doc.add_paragraph("正式提交或队内归档时，应将以下文件与本章节同时保存：sample_index.csv；三模态分阶段 NPZ 与 aligned_50 NPZ；各阶段 extraction_log 和 alignment_log；aligned_50_schema.json；final_sample_validation.csv；official_interface_comparison.json；processing_config.yaml；reproducibility_manifest.json；CPU 与 RTX 4090 环境文件；A–G 阶段源代码。原始视频继续保存在 00_raw_readonly/，不在处理阶段覆盖。")

    for table in doc.tables:
        style_table(table)
    doc.core_properties.title = "问题一求解及分析"
    doc.core_properties.subject = "复杂场景下多模态情感识别的特征建模与时间组织"
    doc.core_properties.comments = "在原始章节草稿基础上补充可核验统计、图表和电子材料索引。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
