from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from docx.oxml.ns import qn

from enrich_q1_docx import configure_styles, set_cell_text, style_table


INPUT = Path(
    r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_26\单个章节\论文-第一问数据处理_模型假设与符号说明稿.docx"
)
OUTPUT = INPUT.with_name("论文-第一问数据处理_模型假设与符号说明_二三问完善稿.docx")


def set_text(paragraph, text: str) -> None:
    paragraph.clear()
    run = paragraph.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    parent = element.getparent()
    if parent is not None:
        parent.remove(element)


def remove_all_rows(table) -> None:
    for row in list(table.rows):
        table._tbl.remove(row._tr)


def fill_table(table, headers: list[str], rows: list[list[str]]) -> None:
    remove_all_rows(table)
    header = table.add_row().cells
    for i, value in enumerate(headers):
        set_cell_text(header[i], value, bold=True, color="FFFFFF", size=9)
    for values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(values):
            set_cell_text(cells[i], value, size=8.8)
    style_table(table)


def insert_body_before(anchor, text: str, lead: str | None = None):
    p = anchor.insert_paragraph_before()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if lead and text.startswith(lead):
        r = p.add_run(lead)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        text = text[len(lead):]
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)
    return p


def insert_assumption_block(anchor, intro: str, assumptions: list[tuple[str, str]]) -> None:
    insert_body_before(anchor, intro, lead="说明。")
    for lead, body in assumptions:
        insert_body_before(anchor, lead + body, lead=lead)


def update_toc(doc: Document) -> None:
    """Add the four newly populated entries to the document's static TOC cache."""
    paragraphs = list(doc.paragraphs)
    toc_q1_notation = next(
        p for p in paragraphs
        if p.style.name.lower() == "toc 2" and p.text.startswith("2.2")
    )
    chapter3 = next(
        p for p in paragraphs
        if p.style.name.lower() == "toc 1" and p.text.startswith("3 ")
    )
    set_text(toc_q1_notation, "2.2 第一问符号说明\t6")
    entries = [
        ("2.3 第二问模型假设", "7"),
        ("2.4 第二问符号说明", "7"),
        ("2.5 第三问模型假设", "8"),
        ("2.6 第三问符号说明", "8"),
    ]
    for title, page in entries:
        p = chapter3.insert_paragraph_before(f"{title}\t{page}", style="toc 2")
        p.paragraph_format.space_after = Pt(0)


def main() -> None:
    doc = Document(str(INPUT))
    configure_styles(doc)
    paragraphs = list(doc.paragraphs)

    q2_heading = next(p for p in paragraphs if p.style.name == "Heading 2" and "第二问模型假设" in p.text)
    q2_notation_heading = next(p for p in paragraphs if p.style.name == "Heading 2" and "第二问符号说明" in p.text)
    q2_placeholder = next(p for p in paragraphs if p.text.startswith("本节待补充"))
    q2_table_caption = next(p for p in paragraphs if "表 2-2" in p.text)
    q2_table_note = next(p for p in paragraphs if p.text.startswith("预留表：待问题二"))

    q3_heading = next(p for p in paragraphs if p.style.name == "Heading 2" and "第三问模型假设" in p.text)
    q3_notation_heading = next(p for p in paragraphs if p.style.name == "Heading 2" and "第三问符号说明" in p.text)
    q3_placeholder = next(p for p in paragraphs if p.text.startswith("本节待补充：可解释性"))
    q3_table_note = next(p for p in paragraphs if p.text.startswith("预留表：待问题三"))

    # Convert the template's placeholder headings into the completed section titles.
    set_text(q2_heading, "2.3 第二问模型假设")
    set_text(q2_notation_heading, "2.4 第二问符号说明")
    set_text(q3_heading, "2.5 第三问模型假设")
    set_text(q3_notation_heading, "2.6 第三问符号说明")

    q2_assumptions = [
        ("假设 1：局部缺失定义为连续位置缺失。", "题目中的缺失是某一模态在连续序列区间内的全部特征维度被置为不可用，不等同于整段样本永久失去该模态，也不等同于随机删除彼此独立的单点。缺失比例按有效内容位置计，不将起始标记、结束标记和 padding 位置计入缺失时长。"),
        ("假设 2：官方数据划分保持不变。", "训练集、验证集和测试集沿用附件 2 给出的划分，不重新随机切分。训练集用于参数学习和标准化统计量估计，验证集用于检查点、模型类型和实验协议选择，测试集只作补充分析；附件 3 没有真实标签，只用于专项推理。"),
        ("假设 3：第二问统一使用 aligned-50 接口。", "模型输入最多包含 50 个序列位置，音频和视觉特征维数分别为 74 和 35；文本使用 text_bert 中的词元编号、注意力标记和分段编号，并由 BERT-Tiny 得到 128 维文本表示。第二问不在训练阶段重新提取问题一的原始视频特征，也不混用 unaligned 版本。"),
        ("假设 4：原始观测 mask 与人工缺失 mask 分离保存。", "以 O 表示数据原生的观测 mask，以 D 表示本次实验构造的人工缺失 mask，最终输入状态由二者逐位置相乘得到。synthetic_missing_mask 只记录人为实验操作，不覆盖 text_mask、audio_mask 和 vision_mask，从而能够区分原始不可观测与实验性遮挡。"),
        ("假设 5：缺失区间只在有效内容位置上构造。", "给定目标缺失比例 r，对满足有效长度条件的样本选择一个或两个连续区间，并保证至少保留一个有效内容位置。训练样本中一部分不新增缺失，其余样本随机选择 1 至 3 个模态和 10%、30%、50%、70% 等目标比例；实际比例因取整及与原生零值重叠可能与目标值略有差异。"),
        ("假设 6：标准化统计量不泄漏验证和测试信息。", "音频与视觉特征的均值、标准差和截断参数只由训练集的原始可用位置估计，验证集、测试集和附件 3/4 复用同一组参数。标准差设置下限以避免近常量维度被放大；不可用位置在标准化后重新置零，词元编号不参与连续数值标准化。"),
        ("假设 7：分类与回归是联合但口径不同的两个任务。", "模型同时输出 Negative、Neutral、Positive 三类情感极性和范围为 [-3, 3] 的连续情感强度。分类使用 Accuracy 和 Macro-F1，回归使用 MAE 和 Pearson 相关系数；若相关系数因预测为常数而不可定义，则保留为空值，不以零代替。"),
        ("假设 8：动态融合权重表示可用信息调节，不表示因果贡献。", "三模态分别投影到共同表示空间，经模态内 Transformer、跨模态注意力和样本级动态融合后进行联合预测。融合权重随观测比例和内容得分变化，只说明模型在当前输入下如何组合信息，不能直接解释为某模态的因果贡献。"),
        ("假设 9：基线与主模型使用可比的训练和评价协议。", "固定融合基线、缺失增强基线、缺失感知主模型及可选重建版本共享数据划分、监督目标和检查点选择规则。重建损失只在相应消融配置中启用，最终 robust_norec 版本不启用重建损失。"),
        ("假设 10：模型选择和结果复现由随机种子及运行记录约束。", "训练代表种子为 42，辅助比较使用 43 和 44；遮挡实验使用固定的遮挡种子。每次运行同时记录模型配置、检查点、标准化参数、数据划分、评价指标和运行设备，避免把不同随机性来源混为一次独立实验。"),
    ]
    insert_assumption_block(q2_notation_heading, "第二问的假设围绕局部连续缺失、训练数据边界、mask 语义、标准化防泄漏和联合预测接口展开。它们只约束问题二的鲁棒预测实验，不把问题三的解释指标提前当作第二问的监督目标。", q2_assumptions)
    remove_paragraph(q2_placeholder)

    q2_symbols = [
        ["样本与划分", "i, N, split_i", "样本编号、样本总数及 train/valid/test 所属划分；划分沿用附件 2。"],
        ["模态", "m∈{T,A,V}", "T、A、V 分别表示文本、音频和视觉。"],
        ["序列位置", "K, k", "K=50；k=0,…,K−1 为统一序列位置。"],
        ["输入特征", "X_i^(T), X_i^(A), X_i^(V)", "第 i 条样本三种模态的 50 位置输入；音频维数 74，视觉维数 35，文本由 text_bert 编码。"],
        ["原始观测 mask", "O_i^(m)", "原始文件中记录的模态可观测状态。"],
        ["人工缺失 mask", "D_i^(m)", "实验构造的连续区间遮挡状态；仅用于缺失实验。"],
        ["最终观测 mask", "M_i^(m)", "O_i^(m)⊙D_i^(m)，表示输入给模型的有效状态。"],
        ["缺失区间", "C_i, L_i, r", "有效位置集合、区间长度和目标缺失比例；缺失时长按位置比例表示。"],
        ["标准化统计量", "μ_m, σ_m, εσ", "由训练集估计的均值、标准差及标准差下限。"],
        ["标准化特征", "X̃_i^(m)", "使用训练集统计量处理后的连续特征；不可用位置重新置零。"],
        ["模态表示", "H_i^(m), q_i^(m)", "模态内编码结果与该模态有效内容位置比例。"],
        ["动态权重", "α_i^(m)", "由内容得分和观测比例共同决定的样本级融合权重。"],
        ["联合表示", "h_i, z_i, u_i", "融合表示、三分类 logits 和连续回归分支的未约束输出。"],
        ["预测结果", "p̂_i, ŷ_i, ĉ_i", "三类概率、情感强度预测值和最大概率对应的类别。"],
        ["监督标签", "y_i^(c), y_i^(r)", "离散类别标签 annotation 与连续强度标签 label。"],
        ["损失函数", "ℓ_cls, ℓ_reg, ℓ_rec, λ", "加权交叉熵、Smooth-L1 回归损失、可选重建损失及联合损失系数。"],
        ["选择分数", "S", "在完整输入及 T/A/V/TAV 30% 缺失条件上综合评价模型的检查点选择分数。"],
        ["随机性", "seed", "训练种子 42、43、44 及遮挡实验固定种子。"],
    ]
    set_text(q2_table_caption, "表 2-2 第二问主要符号说明")
    fill_table(doc.tables[3], ["类别", "符号", "含义"], q2_symbols)
    set_text(q2_table_note, "符号与数据字段对应关系：aligned_50、mask、normalization、missing_scenarios、core_metrics 和 selection 记录。附件 3 无真实标签，因此不为其定义监督评价符号。")

    q3_assumptions = [
        ("假设 1：解释阶段冻结问题二的预测模型。", "第三问不重新训练预测网络，直接加载问题二选定的 robust_norec 检查点、训练集标准化参数和输入接口。解释层只改变输入观测状态，不改变模型结构、模型参数、标准化统计量或分类—强度一致性投影规则。"),
        ("假设 2：完整输入是所有遮挡实验的共同基准。", "对每条样本先保存完整三模态输入的分类概率和情感强度，再分别遮挡文本、音频或视觉模态，或遮挡某一模态的连续窗口。遮挡前后的输出差异用于量化当前模型对输入扰动的敏感度。"),
        ("假设 3：不同模态采用与其输入接口一致的遮挡操作。", "文本在 BERT 编码前将对应 token 替换为 [MASK]，避免完整上下文泄漏到被解释位置；音频和视觉则将窗口内观测 mask 置为不可用，并沿用问题二的缺失表示规则。原始观测 mask 和人为解释 mask 分开保存。"),
        ("假设 4：模态贡献解释为敏感度而非因果效应。", "模态遮挡前后的概率变化、强度变化、绝对贡献和归一化贡献描述冻结模型对输入改变的响应，不等价于真实世界中该模态的因果贡献，也不等价于注意力权重。"),
        ("假设 5：局部证据由连续窗口定义。", "对模态 m 设窗口起点 k、长度 w，则窗口为 [k,k+w)；分类敏感度和回归敏感度分别排序，不能强制两类任务使用相同的 top-k 窗口。正式协议在附件 2 验证集上比较窗口比例后固定为长度 5、步长 2、top-k=3。"),
        ("假设 6：解释协议在专项推理前固定。", "窗口长度、步长、top-k、遮挡种子和输出投影规则均在带标签验证集上确定，不使用附件 4 的预测分布或任何隐含标签反向调参。附件 4 仅用于最终无标签推理和解释文件生成。"),
        ("假设 7：局部位置采用归一化时间语义。", "音频和视觉窗口首先表示为 50 个归一化视频时间位置，只有在获得对应视频时长 D_i 后才转换为秒级区间；文本证据暂以 token 或 bin 位置表示，不能把均匀 token 分桶写成真实词级时间戳。"),
        ("假设 8：无标签专项样本不计算监督性能。", "附件 4 的预测类别、情感强度、模态贡献和局部证据只用于描述模型输出与输入敏感度，不能据此计算 Accuracy、F1、MAE 或 Pearson，也不能把预测分布解释为真实类别分布。"),
        ("假设 9：所有解释记录能够回溯到样本和运行材料。", "每条预测、模态贡献和窗口证据均保留 sample_id、模型检查点、输入版本、窗口边界、有效比例、输出变化和运行元数据，从而能够回到对应的 aligned 文件和原始素材。"),
    ]
    insert_assumption_block(q3_notation_heading, "第三问的假设围绕冻结模型、输入遮挡、窗口证据、归一化位置和无标签推理边界展开。解释结果应回答“当前模型对哪些输入扰动敏感”，不应直接升级为因果结论。", q3_assumptions)
    remove_paragraph(q3_placeholder)

    q3_symbols = [
        ["样本与模态", "i, m∈{T,A,V}", "样本编号与文本、音频、视觉模态。"],
        ["完整输入", "X_i, X_i^(m)", "第 i 条样本的完整三模态输入及其中一个模态输入。"],
        ["冻结模型输出", "p̂_i, ŷ_i, ĉ_i", "完整输入下三类概率、连续强度及预测类别。"],
        ["模态遮挡输入", "X_i^(−m)", "仅遮挡模态 m 后送入冻结模型的输入。"],
        ["遮挡输出", "p̂_i^(−m), ŷ_i^(−m)", "模态 m 被遮挡后的分类概率和强度预测。"],
        ["模态敏感度", "Δ_i,m^cls, Δ_i,m^reg", "遮挡前后分类输出和回归输出的变化量。"],
        ["贡献统计", "A_i,m, R_i,m, C_i,m", "绝对贡献、有符号贡献、归一化贡献及其汇总字段。"],
        ["主要模态", "m_i*", "按规定贡献统计在样本级选出的主要参考模态。"],
        ["窗口", "W_{i,m,k,w}", "模态 m 从位置 k 开始、长度为 w 的连续窗口 [k,k+w)。"],
        ["窗口协议", "w, s, top-k", "窗口长度、滑动步长和保留的关键窗口数；正式值为 5、2、3。"],
        ["窗口输出变化", "δ_i,m,k^cls, δ_i,m,k^reg", "遮挡窗口后分类概率和情感强度的变化。"],
        ["证据位置", "τ_start, τ_end", "归一化时间起止位置；有视频时长时可换算为秒级区间。"],
        ["有效比例", "valid_fraction", "窗口内原始或最终有效观测位置比例。"],
        ["证据排序", "rank_cls, rank_reg", "窗口在分类和回归敏感度排序中的名次。"],
        ["输出文件", "q3_prediction, q3_modal_contribution, q3_temporal_evidence", "分别记录专项预测、模态贡献和局部窗口证据。"],
        ["协议与复现", "protocol, run_metadata, checkpoint_hash", "解释协议、运行元数据和冻结检查点哈希。"],
    ]
    set_text(next(p for p in doc.paragraphs if "表 2-3" in p.text), "表 2-3 第三问主要符号说明")
    fill_table(doc.tables[4], ["类别", "符号", "含义"], q3_symbols)
    set_text(q3_table_note, "符号与数据字段对应关系：q3_prediction.csv、q3_modal_contribution.csv、q3_temporal_evidence.csv、q3_explanation_validation.csv、q3_explanation_protocol.json 和 q3_run_metadata.json。")

    update_toc(doc)
    doc.core_properties.title = "第一问数据处理与三问模型假设符号说明"
    doc.core_properties.subject = "问题一特征处理、问题二鲁棒预测和问题三可解释性模型的假设与符号"
    doc.core_properties.comments = "在第一问模型假设与符号说明稿基础上，根据完整论文补充问题二和问题三。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
