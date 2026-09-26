from __future__ import annotations

from pathlib import Path
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

from enrich_q1_docx import (
    SOURCE,
    WORKSPACE,
    OUTPUT,
    add_caption,
    add_source_note,
    configure_styles,
    set_cell_border,
    set_cell_shading,
    set_cell_text,
    style_table,
    replace_text,
    find_drawing_paragraphs,
    remove_paragraph,
)

OUTPUT = SOURCE.with_name("4问题一求解及分析_正文修订稿.docx")


def paragraph_text(paragraph: object) -> str:
    return getattr(paragraph, "text", "")


def find_paragraph(doc: Document, needle: str):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(f"找不到锚点段落: {needle}")


def insert_picture_before(paragraph, image: Path, caption: str, source: str, width: float = 6.15) -> None:
    p = paragraph.insert_paragraph_before()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(image), width=Inches(width))
    cap = paragraph.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)
    note = paragraph.insert_paragraph_before(source)
    note.paragraph_format.space_after = Pt(5)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(8)


def insert_table_before(paragraph, caption: str, headers: list[str], rows: list[list[str]], source: str):
    cap = paragraph.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)

    table = paragraph._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    paragraph._p.addprevious(table._element)
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, color="FFFFFF")
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            set_cell_text(cells[j], value)
    style_table(table)
    note = paragraph.insert_paragraph_before(source)
    note.paragraph_format.space_after = Pt(5)
    for run in note.runs:
        run.italic = True
        run.font.size = Pt(8)
    return table


def omml_text(text: str):
    omath = OxmlElement("m:oMathPara")
    math = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text_node = OxmlElement("m:t")
    text_node.text = text
    run.append(text_node)
    math.append(run)
    omath.append(math)
    return omath


def replace_paragraph_with_equation(paragraph, equation: str) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)
    p.append(omml_text(equation))
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(4)


def convert_formulas(doc: Document) -> None:
    replacements = {
        "sᵢ = strip(video_idᵢ) ‖ \"_\" ‖ strip(clip_idᵢ)": "s_i = strip(video_id_i) + \"_\" + strip(clip_id_i)",
        "b(j) = min(⌊jK/Mᵢ⌋, K−1)": "b(j) = min(floor(j K / M_i), K - 1)",
        "xᵀᵢₖ = (1/|Jᵢₖ|) Σⱼ∈Jᵢₖ hᵢⱼ    mᵀᵢₖ = 1(|Jᵢₖ|>0)": "x^T_ik = (1 / |J_ik|) sum_(j in J_ik) h_ij       m^T_ik = 1(|J_ik| > 0)",
        "aᵢᵣ = [MFCC₁₃, ΔMFCC₁₃, Δ²MFCC₁₃, q₄, s₄, LogMel₂₇]": "a_ir = [MFCC_13, DeltaMFCC_13, Delta^2MFCC_13, q_4, s_4, LogMel_27]",
        "vᵢq = [p₃₀, c₄, g]": "v_iq = [p_30, c_4, g]",
        "Iᵢₖ = [kDᵢ/K, (k+1)Dᵢ/K)    k=0,1,…,49": "I_ik = [k D_i / K, (k + 1) D_i / K),    k = 0, 1, ..., 49",
        "Xᵢₖ^(m) = mean({z : time(z)∈Iᵢₖ})    Mᵢₖ^(m) = 1(nᵢₖ^(m)>0)": "X_ik^(m) = mean({z : time(z) in I_ik})       M_ik^(m) = 1(n_ik^(m) > 0)",
    }
    for p in doc.paragraphs:
        if p.text in replacements:
            replace_paragraph_with_equation(p, replacements[p.text])


def main() -> None:
    doc = Document(str(SOURCE))
    configure_styles(doc)

    replace_text(doc, "问题一求解及分析章节草稿", "问题一求解及分析")
    replace_text(doc, "本章草稿只覆盖问题一。它把队友的工程记录重组为“输入—表示—对齐—输出—核验”的论文逻辑，不包含问题二和问题三的预测模型，也不把特征覆盖率当作分类或回归性能。", "本章围绕问题一建立文本、音频和视觉三模态特征提取与时序组织流程。章节依次说明样本索引、模态特征表示、统一 50 桶接口、质量核验和复现记录，并将每个结论对应到可追溯的数据文件和图表。特征覆盖率用于描述输入质量，不作为分类或回归性能指标。")
    replace_text(doc, "现有 text_index.csv 可直接核验 100 条文本索引，text_feature_validation.csv 中 100 条记录均标记为 PASS。队友运行笔记进一步记录：标签行数、MP4 数和唯一 sample_id 数均为 100，缺失视频和歧义匹配均为 0；这些汇总值在正式提交前仍需与原始 sample_index.csv 复算。", "工作区复核得到：标签行数、MP4 数和唯一 sample_id 数均为 100，缺失视频和歧义匹配均为 0，视频元信息可读率为 100/100；text_feature_validation.csv 中 100 条记录均为 PASS。")
    replace_text(doc, "表 4-5 三模态位置覆盖汇总 统计值待原始结果文件复算", "表 4-5 三模态位置覆盖汇总。完整统计来自 multimodal_coverage_summary.csv。")
    replace_text(doc, "图 4-3 三模态位置覆盖与样本级可观测性 统计值待原始结果文件复算", "图 4-3 三模态总体质量概览。图中同时展示样本级有效桶分布、总体覆盖率以及至少一个有效桶和 50 桶全有效样本数。")
    replace_text(doc, "表 4-6 三类边界样本的桶级回溯摘要 统计值待原始结果文件复算", "表 4-6 三类边界样本的桶级回溯摘要。完整记录来自 representative_traceability.csv。")
    replace_text(doc, "且本次收到的材料缺少音频、视觉和最终验收原始文件。", "当前工作区已经保存音频、视觉和最终验收文件；正式提交时仍应将完整特征文件、日志、配置和代码按赛事要求一并归档。")
    replace_text(doc, "该接口可以直接用于问题二的局部模态缺失建模和问题三的证据定位。", "该接口可以直接作为问题二的输入。若后续构造人为缺失场景，应新增 synthetic_missing_mask，不能覆盖问题一记录的原始模态 mask。")

    # Remove the old end-of-section figure placeholder after the real figure is inserted inline.
    for p in list(doc.paragraphs):
        if p.text.startswith("图 4-3 三模态总体质量概览。图中同时展示"):
            remove_paragraph(p)

    if doc.tables:
        set_cell_text(doc.tables[0].cell(0, 0), "核验说明  本修订稿使用当前 E_workspace 中的 sample_index、三模态特征、aligned-50、质量核验和复现文件生成。所有统计数字均可由对应 CSV、JSON 或 NPZ 文件复核。")
        set_cell_text(doc.tables[3].cell(0, 0), "实现口径  视觉主体选择采用实际运行规则：每个采样帧独立选择 landmark 外接框面积最大的人脸。检测不到人脸时不沿用上一帧特征，并将 vision_mask 置为 0。")

    # Remove the old third image, which was a low-information placeholder.
    drawings = find_drawing_paragraphs(doc)
    if len(drawings) >= 3:
        remove_paragraph(drawings[2])

    # Convert standalone formula paragraphs before inserting evidence figures.
    convert_formulas(doc)

    # Inline figures at the point where their results are discussed.
    fig1 = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig1_multimodal_qc_overview/fig1_multimodal_qc_overview.png"
    fig2 = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig2_mask_structure_revised/fig2_mask_structure_revised.png"
    fig3 = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig3_alignment_profiles/fig3_alignment_profiles.png"
    fig4 = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig4_label_stratified_availability/fig4_label_stratified_availability.png"
    fig6 = WORKSPACE / "02_features_q1/stepE_alignment/image/vpublication_chinese/fig6_pipeline_traceability/fig6_pipeline_traceability.png"
    fig7 = WORKSPACE / "02_features_q1/stepF_validation/image/publication_chinese/fig7_final_validation_audit/fig7_final_validation_audit.png"
    fig8 = WORKSPACE / "02_features_q1/stepG_repro/image/publication_chinese/fig8_reproducibility_record_audit/fig8_reproducibility_record_audit.png"
    representative = WORKSPACE / "02_features_q1/stepF_validation/image/representative_traceability.png"
    text_token = WORKSPACE / "02_features_q1/stepB_text/image/text_token_count_distribution.png"
    text_by_label = WORKSPACE / "02_features_q1/stepB_text/image/text_token_count_by_annotation.png"
    text_mask = WORKSPACE / "02_features_q1/stepB_text/image/text_mask_coverage.png"

    insert_picture_before(find_paragraph(doc, "为兼顾官方接口兼容性"), fig1, "图 4-3 三模态总体质量概览", "数据来源：stepE_alignment/data/multimodal_coverage_summary.csv；文本桶为 token 序列位置，音频和视觉桶为视频相对时间位置。")
    insert_picture_before(find_paragraph(doc, "空位置填充零向量"), text_token, "图 4-4 文本内容 token 数分布", "数据来源：stepB_text/data/text_index.csv；中位数为 21。")
    insert_picture_before(find_paragraph(doc, "空位置填充零向量"), text_by_label, "图 4-5 按情感标签分层的文本 token 数", "数据来源：stepB_text/data/text_index.csv；该图为描述性统计。")
    insert_picture_before(find_paragraph(doc, "空位置填充零向量"), text_mask, "图 4-6 文本 50 个 token 序列位置的观测 mask", "数据来源：stepB_text/image/text_mask_coverage.png；空桶不是视频时间缺失。")
    insert_picture_before(find_paragraph(doc, "零向量只是填充值"), fig2, "图 4-7 按正确序列语义显示的三模态观测 mask", "数据来源：stepE_alignment/data/multimodal_sample_audit.csv；文本轴不是视频时间轴。")
    insert_picture_before(find_paragraph(doc, "零向量只是填充值"), fig3, "图 4-8 三模态逐桶可用率与音视频联合 mask", "数据来源：stepE_alignment/data/multimodal_alignment_log.csv。")
    insert_picture_before(find_paragraph(doc, "文本的 45.18%"), fig4, "图 4-9 按情感标签分层的三模态有效桶数", "数据来源：stepE_alignment/data/multimodal_sample_audit.csv；仅用于检查可观测性分布。")
    insert_picture_before(find_paragraph(doc, "文本的 45.18%"), fig6, "图 4-10 样本在 Step A–E 流程中的保留情况", "数据来源：stepG_repro/data/unified_extraction_log.csv；失败样本保留并由状态和 mask 标记。")
    insert_picture_before(find_paragraph(doc, "质量检查覆盖 dtype"), fig7, "图 4-11 aligned-50 特征集最终验收审计", "数据来源：stepF_validation/data/final_sample_validation.csv、numeric_quality_summary.csv 和 official_interface_comparison.json。")
    insert_picture_before(find_paragraph(doc, "上述回溯说明"), representative, "图 4-12 代表性样本的 50 桶回溯", "数据来源：stepF_validation/data/representative_traceability.csv；该图证明位置可回溯，不证明文本桶具有真实词级时间戳。")
    insert_picture_before(find_paragraph(doc, "运行笔记记录随机种子"), fig8, "图 4-13 问题一流程的复现记录覆盖情况", "数据来源：stepG_repro/data/reproducibility_manifest.json、unified_extraction_log.csv 和 processing_config.yaml。")

    # Put high-value numeric tables at the corresponding result paragraphs.
    insert_table_before(find_paragraph(doc, "文本的 45.18%"), "表 4-7 三模态覆盖汇总", ["模态", "有效桶", "可能桶", "覆盖率", "至少一个有效桶", "50 桶全有效"], [
        ["文本", "2259", "5000", "45.18%", "100", "6"],
        ["音频", "4957", "5000", "99.14%", "100", "63"],
        ["视觉", "3865", "5000", "77.30%", "90", "47"],
    ], "数据来源：stepE_alignment/data/multimodal_coverage_summary.csv。文本覆盖率表示 token 序列位置占用率，音频和视觉覆盖率表示视频时间桶有效观测率。")
    insert_table_before(find_paragraph(doc, "质量检查覆盖 dtype"), "表 4-8 数值质量核验摘要", ["模态", "dtype", "有效值比例", "mask 非法样本", "全零行", "mask=1 且全零行", "相邻有效桶 L2 中位数"], [
        ["文本", "float32", "100%", "0", "2741", "0", "14.0345"],
        ["音频", "float32", "100%", "0", "43", "0", "1274.8173"],
        ["视觉", "float32", "100%", "0", "1135", "0", "0.0812"],
    ], "数据来源：stepF_validation/data/numeric_quality_summary.csv。全零行主要位于 mask=0 的不可观测位置。")
    insert_table_before(find_paragraph(doc, "官方附件 2 的 aligned_50.pkl"), "表 4-9 官方接口交叉核验", ["项目", "结果"], [
        ["官方样本总数", "4850"],
        ["train / valid / test", "3395 / 728 / 727"],
        ["text / audio / vision 形状", "N×50×768 / N×50×74 / N×50×35"],
        ["共享 sample_id", "18"],
        ["分类标签不一致数", "0"],
        ["回归标签不一致数", "0"],
    ], "数据来源：stepF_validation/data/official_interface_comparison.json。该核验比较接口、ID 和标签口径，不比较特征逐元素数值。")

    # Keep the reproducibility table close to its discussion, not at the document tail.
    insert_table_before(find_paragraph(doc, "运行笔记记录随机种子"), "表 4-10 复现材料清单", ["材料", "文件", "作用"], [
        ["处理配置", "processing_config.yaml", "固定维度、时间桶、mask 和随机种子"],
        ["统一日志", "unified_extraction_log.csv", "记录 A–E 阶段逐样本状态"],
        ["复现清单", "reproducibility_manifest.json", "记录关键文件和代码 SHA-256"],
        ["环境", "environment_cpu.yaml / environment_cuda4090.yaml", "记录 CPU 与 RTX 4090 环境"],
        ["源代码", "stepA–stepG/code/*.py", "固定处理算法与运行顺序"],
    ], "数据来源：stepG_repro/data 和 01_manifest/env。当前清单尚未登记全部单样本 NPZ 的逐文件哈希。")

    # The old end appendix from the previous revision is not present because we start from SOURCE.
    for table in doc.tables:
        style_table(table)
    doc.core_properties.title = "问题一求解及分析"
    doc.core_properties.subject = "复杂场景下多模态情感识别的特征建模与时间组织"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
