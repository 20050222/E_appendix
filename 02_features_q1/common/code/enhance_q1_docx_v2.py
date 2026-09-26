from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from enrich_q1_docx import configure_styles, set_cell_text, style_table


INPUT = Path(r"E:\备份\读研相关\2026-2027研一\04_竞赛与实践\数学建模\E题建模\论文\2026_09_25\4问题一求解及分析_正文扩展稿.docx")
OUTPUT = INPUT.with_name("4问题一求解及分析_正文增强稿.docx")


def find_paragraph(doc: Document, needle: str):
    for paragraph in doc.paragraphs:
        if needle in paragraph.text:
            return paragraph
    raise ValueError(f"找不到锚点段落: {needle}")


def add_body_before(paragraph, text: str, lead: str | None = None) -> None:
    p = paragraph.insert_paragraph_before()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if lead and text.startswith(lead):
        run = p.add_run(lead)
        run.bold = True
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(10.5)
        text = text[len(lead):]
    run = p.add_run(text)
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(10.5)


def add_heading_before(paragraph, text: str, level: int = 3):
    p = paragraph.insert_paragraph_before(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    return p


def add_table_before(paragraph, caption: str, headers: list[str], rows: list[list[str]], source: str):
    cap = paragraph.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    run = cap.add_run(caption)
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(9.5)

    table = paragraph._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    paragraph._p.addprevious(table._element)
    for i, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color="FFFFFF")
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value)
    style_table(table)
    set_table_pagination(table)

    note = paragraph.insert_paragraph_before(source)
    note.paragraph_format.space_after = Pt(6)
    for item in note.runs:
        item.italic = True
        item.font.name = "Microsoft YaHei"
        item._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        item.font.size = Pt(8)
    return table


def set_table_pagination(table) -> None:
    """Repeat table headers and prevent individual rows from splitting."""
    if not table.rows:
        return
    header_trpr = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    header_trpr.append(header)
    for row in table.rows:
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:cantSplit"))


def renumber_captions(doc: Document) -> None:
    figure_no = 0
    table_no = 0
    pattern = re.compile(r"^(图|表)\s*4-(?:\d+|XX)")
    for paragraph in doc.paragraphs:
        match = pattern.match(paragraph.text.strip())
        if not match:
            continue
        kind = match.group(1)
        if kind == "图":
            figure_no += 1
            number = figure_no
        else:
            table_no += 1
            number = table_no
        old_prefix = match.group(0)
        new_prefix = f"{kind} 4-{number}"
        replacement = paragraph.text.strip().replace(old_prefix, new_prefix, 1)
        for run in paragraph.runs:
            run.text = ""
        if paragraph.runs:
            paragraph.runs[0].text = replacement
        else:
            paragraph.add_run(replacement)


def main() -> None:
    doc = Document(str(INPUT))
    configure_styles(doc)

    # 4.1: make the acceptance logic explicit before the data-engineering details.
    anchor_42_content = find_paragraph(doc, "首先以标签表为主表建立 100 条样本索引")
    add_body_before(anchor_42_content, "本题的验收对象不是一个孤立的特征矩阵，而是一条可以从原始视频回溯到最终特征的证据链。完整性对应样本数量、sample_id 唯一性和模态文件一一对应；可核验性对应位置编号、有效长度、时间边界、源观测数和 mask；可复现性对应模型版本、关键参数、环境文件、代码哈希和逐样本日志。只有三类证据同时存在，才能把‘提取出了特征’写成可审计的建模结果。", lead="验收逻辑。")
    add_body_before(anchor_42_content, "Step A 至 Step G 的关系是串联而非并列：Step A 固定主键，Step B–D 在同一主键下生成三种模态，Step E 统一序列接口，Step F 检查数值与官方字段口径，Step G 固定复现材料。后续任何统计图都应能回到其中一个 CSV 或配置文件，而不把图中的汇总数字作为脱离原始记录的独立结论。", lead="阶段衔接。")
    add_table_before(anchor_42_content, "表 4-XX 题目要求与可核验证据的对应关系", ["题目要求", "判定指标", "证据文件", "本次结果"], [
        ["原始样本覆盖完整", "标签行数、MP4 数、唯一 sample_id、缺失路径", "stepA_index/data/sample_index.csv", "100 / 100 / 100 / 0"],
        ["特征一一对应", "三模态 NPZ 和 aligned 文件是否按 sample_id 存在", "各阶段 extraction_log、multimodal_sample_audit.csv", "100 条索引均保留"],
        ["时序可核验", "bin_index、边界、source_count、valid/mask", "multimodal_alignment_log.csv、audio_alignment_log.csv、vision_alignment_log.csv", "50 个位置均有接口记录"],
        ["方法可复现", "模型、参数、环境、代码和状态日志", "processing_config.yaml、环境 YAML、reproducibility_manifest.json", "A–E 共 500 条样本级状态记录"],
    ], "数据来源：stepA–stepG/data、docs 和 01_manifest/env；数值均来自当前工作区文件。")

    # 4.2: explain what the index and status fields mean, not only their counts.
    anchor_43 = find_paragraph(doc, "4.3 三模态特征表示模型")
    add_body_before(anchor_43, "sample_index.csv 的作用是把‘样本是否存在’与‘某个模态是否可观测’分开。index_status=PASS 只说明原始视频已经被定位并能读取元信息，不等于文本、音频或视觉三种特征都一定有效；后者由各模态的 status、valid_bins 和 mask 共同判断。这样的分层可以避免因为视觉检测失败而删除整条样本，也避免把文件存在误写为模态质量通过。", lead="状态字段的边界。")

    # 4.3: add interpretation paragraphs immediately before each next modality.
    anchor_audio = find_paragraph(doc, "4.3.2 音频特征表示")
    add_body_before(anchor_audio, "图表解读。文本 token 数分布集中在较短区间，说明 50 个位置主要是为统一接口服务，而不是表示每条转写都包含 50 个独立语义片段。按 annotation 分层的图没有显示出明显的长度分层结论，因此后续模型不应把 token 数直接当作情感强度；它更适合作为文本可观测性和截断风险的质量指标。文本 mask 图中的后部空位是序列填充，而不是视频尾部没有声音或画面。", lead="文本阶段结果。")
    anchor_vision = find_paragraph(doc, "4.3.3 视觉特征表示")
    add_body_before(anchor_vision, "音频结果的重点是‘时间覆盖稳定、尾部存在轻微解码差异’。99.14% 的有效桶率和 47–50 的有效桶范围表明，音频通常覆盖整段视频；约 -0.10 秒的时长差主要影响末端桶，不能据此认为中间时间段发生了系统性错位。图中的音频诊断曲线应与 audio_mask 一起阅读，空桶的零向量不参与有效均值。", lead="音频阶段结果。")
    anchor_44 = find_paragraph(doc, "4.4 统 一时间轴与多模态对齐")
    add_body_before(anchor_44, "视觉结果揭示了本题最需要保留的风险信息：10 条样本没有检测到有效人脸，但这些样本仍有完整的 sample_id、音频和文本记录。视觉覆盖率 77.30% 不是算法准确率，而是在人脸检测规则下的可观测桶比例；因此论文中应同时报告 PASS/FAIL、失败原因和 vision_mask，不能只报告一个平均覆盖率。多脸样本按最大脸框选主体的规则被写入日志，保证同一版本重跑时选择结果确定。", lead="视觉阶段结果。")

    # 4.4: make the semantic difference between common length and common time explicit.
    anchor_45 = find_paragraph(doc, "4.5 输出文件组织与可复现设计")
    add_body_before(anchor_45, "对齐后的 50 个位置具有统一的数组下标，但不应把三种模态的第 k 位解释成完全相同的物理事件。音频和视觉第 k 位对应视频时长的归一化区间；文本第 k 位对应 token 顺序的近似位置。该差异已经通过 alignment_interpretation 字段写入 CSV，后续融合模型可以共享长度并使用独立 mask，同时在论文中如实说明文本没有词级时间戳。", lead="统一接口的解释边界。")
    add_table_before(anchor_45, "表 4-XX 三模态特征文件的字段级接口", ["模态", "特征数组", "mask 与有效长度", "额外可追溯字段", "缺失时的处理"], [
        ["文本", "text_features: 50×768", "text_mask、text_valid_length", "token_index_map、token_count、tokenizer_version", "空桶置零，mask=0；不声明为真实时间缺失"],
        ["音频", "audio_features: 50×74", "audio_mask、valid_bins", "bin_start/end_sec、source_frame_count、sample_rate", "无有效帧的桶置零，mask=0"],
        ["视觉", "vision_features: 50×35", "vision_mask、valid_bins", "face_count、selected_face、source_frame_count", "无有效人脸的桶置零，记录失败原因"],
    ], "数据来源：stepB_text、stepC_audio、stepD_vision 和 stepE_alignment 的 schema、日志及 NPZ 输出。")

    # 4.6: strengthen the interpretation of the coverage and validation tables.
    anchor_461 = find_paragraph(doc, "根据队友的最终运行记录")
    add_body_before(anchor_461, "覆盖率的比较必须带上分母。文本的 2259/5000 是 100 条样本在 50 个 token 位置上的非空比例；音频的 4957/5000 和视觉的 3865/5000 是视频时间桶中存在有效源帧的比例。三者数值可以并列展示，但不能直接横向排序为‘音频优于文本、视觉劣于文本’，因为文本的桶不是视频时间桶。论文表述应使用‘位置占用率’和‘时间桶可观测率’这两个不同名称。", lead="覆盖率的正确解释。")
    anchor_47 = find_paragraph(doc, "4.7 本问小结")
    add_body_before(anchor_47, "从题目要求逐项核对，第一问已经形成四类可提交材料：样本索引材料回答‘处理了哪些样本’，三模态特征与 mask 回答‘提取了什么’，对齐日志回答‘如何组织到 50 个位置’，质量与复现材料回答‘结果能否被复核和重跑’。其中 90/100 三模态完整并不意味着剩余 10 条应删除，而是说明后续建模需要使用 mask 感知的融合策略，或在实验中单独报告视觉缺失子集。", lead="题目要求对照。")
    add_table_before(anchor_47, "表 4-XX 第一问可交付文件清单", ["文件类别", "核心文件", "正文中的说明", "提交注意事项"], [
        ["索引与标签", "sample_index.csv、label_summary.csv、label_consistency_report.md", "样本主键、标签和路径回溯", "保留原始行号、状态和失败原因"],
        ["单模态特征", "text/audio/vision NPZ 与 extraction_log.csv", "说明特征定义、维度和参数", "NPZ 与日志必须按 sample_id 一一对应"],
        ["对齐接口", "aligned_50、multimodal_alignment_log.csv、schema JSON", "说明 50 桶、边界、聚合和 mask", "区分 token 序列位置与视频时间位置"],
        ["质量与复现", "final_sample_validation.csv、numeric_quality_summary.csv、环境 YAML、配置和代码", "说明数值合法性及重跑条件", "原始数据只读，处理结果写入其他目录"],
    ], "文件位置：E_workspace/01_manifest、02_features_q1/stepA–stepG；正文只展示汇总，完整文件作为电子材料归档。")

    # Captions had been inserted in several passes. Normalize their local sequence.
    renumber_captions(doc)
    for table in doc.tables:
        # Existing and newly inserted tables share the same cross-page behavior.
        if not table.rows or table.rows[0]._tr.xpath("./w:trPr/w:tblHeader"):
            continue
        set_table_pagination(table)
    doc.core_properties.title = "问题一求解及分析"
    doc.core_properties.subject = "复杂场景下多模态情感识别的特征建模与时间组织"
    doc.core_properties.comments = "在正文扩展稿基础上补充验收逻辑、CSV 字段解释、真实统计含义和交付文件映射。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
