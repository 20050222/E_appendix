from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from enrich_q1_docx import (
    SOURCE,
    WORKSPACE,
    set_cell_text,
    style_table,
    configure_styles,
    replace_text,
)


CURRENT = SOURCE.with_name("4问题一求解及分析_正文修订稿.docx")
OUTPUT = SOURCE.with_name("4问题一求解及分析_正文扩展稿.docx")


def find_paragraph(doc: Document, needle: str):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise ValueError(f"找不到锚点段落: {needle}")


def insert_text_before(paragraph, text: str, *, lead: str | None = None) -> None:
    p = paragraph.insert_paragraph_before()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    if lead and text.startswith(lead):
        r = p.add_run(lead)
        r.bold = True
        r.font.name = "Microsoft YaHei"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        r.font.size = Pt(10.5)
        text = text[len(lead):]
    r = p.add_run(text)
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(10.5)


def insert_table_before(paragraph, caption: str, headers: list[str], rows: list[list[str]], source: str):
    cap = paragraph.insert_paragraph_before()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.keep_with_next = True
    r = cap.add_run(caption)
    r.bold = True
    r.font.name = "Microsoft YaHei"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    r.font.size = Pt(9.5)

    table = paragraph._parent.add_table(rows=1, cols=len(headers), width=Inches(6.5))
    paragraph._p.addprevious(table._element)
    for j, value in enumerate(headers):
        set_cell_text(table.rows[0].cells[j], value, bold=True, color="FFFFFF")
    for values in rows:
        cells = table.add_row().cells
        for j, value in enumerate(values):
            set_cell_text(cells[j], value)
    style_table(table)
    note = paragraph.insert_paragraph_before(source)
    note.paragraph_format.space_after = Pt(6)
    for run in note.runs:
        run.italic = True
        run.font.name = "Microsoft YaHei"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        run.font.size = Pt(8)
    return table


def main() -> None:
    doc = Document(str(CURRENT))
    configure_styles(doc)

    # The previous revision already contains figures at their relevant locations.
    # This pass adds the explanatory prose and compact CSV-derived result tables beside them.
    p42 = find_paragraph(doc, "4.2 数据索引与预处理")
    insert_text_before(p42, "索引阶段承担的是样本级主键管理，而不是特征提取。标签表中的每一行先被转换为一个 sample_id，再通过 video_relative_path 回到附件 1 的原始 MP4。该顺序保证后续文本、音频和视觉结果都能回到同一条标签记录；即使某个模态提取失败，样本仍然保留在统一索引中。", lead="索引逻辑。")
    insert_text_before(p42, "sample_index.csv 同时保存文件字节数和 SHA-256，因而可以在处理前确认输入文件没有被替换。duration_sec、fps 和 frame_count 用于定义后续视频时间桶及抽帧策略；source_label_row 则保留样本在 label-100.xlsx 中的原始行号，便于从最终特征反向定位标签来源。", lead="字段作用。")

    p43 = find_paragraph(doc, "4.3 三模态特征表示模型")
    insert_table_before(p43, "表 4-11 Step A 样本索引 CSV 字段分组", ["字段组", "代表字段", "论文中的作用"], [
        ["主键与路径", "sample_id、video_id、clip_id、video_relative_path", "建立样本唯一性和原始视频回溯关系"],
        ["媒体元信息", "duration_sec、fps、frame_count、frame_count_source", "定义视频相对时间桶和抽帧记录口径"],
        ["标签输入", "raw_text、label、annotation、source_label_row", "保存文本、连续标签、离散标签和标签表行号"],
        ["完整性与校验", "video_file_size、video_sha256、index_status、failure_reason", "判断输入是否可读取以及失败原因"],
    ], "数据来源：stepA_index/data/sample_index.csv；100 行全部为 PASS。")

    p431 = find_paragraph(doc, "4.3.1 文本特征表示")
    insert_text_before(p431, "文本 CSV 采用“索引—提取日志—特征校验”三级记录。text_index.csv 保存 raw_text 与 normalized_text 的并列版本，说明规范化没有覆盖原始转写；text_extraction_log.csv 记录每条样本的 token_count、实际输入长度、设备和模型；text_feature_validation.csv 则检查 NPZ 是否存在、形状是否为 (50, 768) 以及是否有异常数值。", lead="CSV 记录。")
    insert_text_before(p431, "实测 100 条样本的内容 token 数范围为 5–82，中位数为 21，均值为 23.17，最大值为 82；original_token_count 与 token_count 保持一致，truncated=False 的样本数为 100。因而本次文本结果没有发生最大长度截断，短文本产生的空桶属于均匀分桶后的结构性空位。", lead="文本统计。")
    insert_table_before(p431, "表 4-12 文本特征提取日志的字段与实测结果", ["字段", "含义", "本次结果"], [
        ["text_status", "文本规范化和编码状态", "PASS：100/100"],
        ["token_count", "去除特殊 token 后的内容 token 数", "5–82；中位数 21；均值 23.17"],
        ["truncated", "是否超过 max_length=256", "False：100/100"],
        ["text_valid_length", "50 个序列桶中非空桶数", "与 text_mask 的有效位置数一致"],
        ["feature_shape / device", "输出形状和运行设备", "(50,768)；CPU 运行记录"],
    ], "数据来源：stepB_text/data/text_extraction_log.csv 和 text_feature_validation.csv。")

    p432 = find_paragraph(doc, "4.3.2 音频特征表示")
    insert_text_before(p432, "音频首先通过 FFmpeg 解码为 16 kHz、单声道的 float32 波形，并采用 25 ms 帧长和 10 ms 帧移进行分帧，边界按照左闭右开区间处理。每个音频帧构造 74 维特征向量：其中 39 维来自 13 个 MFCC 及其一阶、二阶差分，4 维为低阶时域或韵律特征，4 维为频谱特征，其余 27 维为 Log-Mel 能量。帧中心时间用于确定其所属的视频相对时间桶，桶内有效帧取均值，从而得到 audio_features∈R^(50×74)、audio_mask、每桶源帧数及对应时间边界。", lead="音频特征。")
    insert_text_before(p432, "音频日志同时保存视频时长、解码音频时长、frame_count、sample_rate、frame_ms、hop_ms、feature_schema 和 FFmpeg 路径，便于复核输入设置及尾部时长差异。本次音频特征的有效桶数范围为 47–50，中位数为 50；解码音频时长与视频时长之差的均值为 -0.1038 秒，中位数为 -0.0992 秒，范围为 -0.1831–0.0026 秒。没有有效音频帧的桶填零并令 audio_mask=0，不用静音值冒充有效观测。", lead="音频核验。")
    insert_table_before(p432, "表 4-13 音频提取日志的质量字段", ["字段", "含义", "本次汇总"], [
        ["audio_video_duration_diff_sec", "解码音频时长减视频时长", "均值 -0.1038 s；范围 -0.1831–0.0026 s"],
        ["valid_bins / mask_ratio", "有效时间桶数及比例", "47–50；中位数 50/50"],
        ["frame_count", "解码后帧级特征数量", "随视频时长变化，逐样本保留"],
        ["sample_rate / frame_ms / hop_ms", "采样和帧级参数", "16000 Hz / 25 ms / 10 ms"],
        ["feature_schema", "固定列顺序和特征定义", "audio74_v1"],
    ], "数据来源：stepC_audio/data/audio_extraction_log.csv、audio_alignment_log.csv 和 audio_feature_schema.json。")

    p433 = find_paragraph(doc, "4.3.3 视觉特征表示")
    insert_text_before(p433, "视觉阶段沿用同一套 50 个视频相对时间桶，并保留视频级 extraction_log、帧级 vision_frame_log 和桶级 vision_alignment_log。视频级日志记录采样帧数、有效人脸帧数、有效桶数和整体状态；帧级日志记录每帧检测到的人脸数量、主体选择结果及失败状态；桶级日志记录采样帧到 50 个时间区间的映射。三层日志共同保证单帧检测失败不会被误判为整段视频缺少视觉信息。", lead="视觉日志。")
    insert_text_before(p433, "实测采样帧数范围为 23–293，中位数为 67.5；检测到有效人脸的帧数范围为 0–283，中位数为 59.5。视觉提取日志中有 90 条样本通过检查，10 条样本因未检测到有效人脸帧而标记为 FAIL(no_valid_face_frames)。此外，仅有 1 条样本出现多脸帧，最多有 22 个采样帧检测到多张人脸，实际处理时按照人脸框面积选择主体。最终，视觉缺失样本继续保留在对齐数据中，并通过 vision_mask 区分有效观测与不可观测区间。", lead="视觉统计。")
    insert_table_before(p433, "表 4-14 视觉提取日志的字段与失败状态", ["字段", "含义", "本次结果"], [
        ["sampled_fps / sampled_frames", "抽帧频率和抽样帧数", "约 10 fps；23–293 帧"],
        ["face_detected_frames", "检测到有效人脸的帧数", "0–283 帧；中位数 59.5"],
        ["valid_bins / mask_ratio", "有效视觉时间桶数", "0–50；中位数 49"],
        ["multiple_face_samples", "检测到多脸的采样帧数", "仅 1 条样本非零，最大 22"],
        ["timestamp_source", "帧时间戳来源", "记录 frame_index_fps_fallback 等实际来源"],
        ["error", "视觉失败原因", "10 条样本为 no_valid_face_frames"],
    ], "数据来源：stepD_vision/data/vision_extraction_log.csv 和 vision_frame_log.csv。")

    p44 = find_paragraph(doc, "4.4 统 一时间轴与多模态对齐")
    insert_text_before(p44, "对齐阶段并不重新计算模态特征，而是读取 Step B–D 的单模态 NPZ，并依据 sample_id 和每条视频的 duration_sec 生成统一接口。每个桶都保存 bin_index、bin_start_sec、bin_end_sec、source_count、valid 和 aggregation；因此读者可以从 aligned 文件的第 k 个位置追溯到该位置由多少个 token、音频帧或有效脸帧聚合而来。", lead="对齐输入。")
    insert_text_before(p44, "multimodal_alignment_log.csv 还显式记录文本桶的解释为 approximate token sequence bin; not timestamp，音视频桶的解释为 normalized video-time bin。这个字段是论文中区分“统一数组长度”和“真实时间同步”的直接证据。", lead="语义标记。")
    insert_table_before(p44, "表 4-15 对齐日志字段与 mask 语义", ["字段", "含义", "论文解读"], [
        ["bin_index", "0–49 的桶编号", "对应 aligned-50 中的序列位置"],
        ["bin_start_sec / bin_end_sec", "视频相对时间边界", "只对音频和视觉具有真实视频时间含义"],
        ["source_count", "桶内源 token 或源帧数量", "解释桶内均值由多少观测构成"],
        ["valid", "桶是否有有效观测", "等价于该模态在该桶的原始 mask"],
        ["aggregation / alignment_interpretation", "聚合方式和位置解释", "区分 mean、mean_valid_faces 与 token sequence bin"],
    ], "数据来源：stepE_alignment/data/multimodal_alignment_log.csv、audio_alignment_log.csv 和 vision_alignment_log.csv。")

    p46 = find_paragraph(doc, "4.6 结果与质量核验")
    insert_text_before(p46, "结果解读需要区分三种层次。第一层是文件层面的完整性，即 100 个 sample_id 是否都有 aligned 文件；第二层是结构层面的合法性，即数组形状、dtype、有限值和 mask 是否符合 schema；第三层是观测层面的完整性，即一个样本的三个模态是否都至少有一个有效桶。当前三层结果分别为 100/100 文件和结构通过、90/100 三模态完整，10 条样本属于文件结构完整但视觉不可观测。", lead="验收层次。")

    p462 = find_paragraph(doc, "4.6.2 数值质量与边界样本回溯")
    insert_text_before(p462, "final_sample_validation.csv 以样本为行，列出每个模态的 dtype、有效桶数、全零行数、mask=1 且全零行数以及相邻有效桶的 L2 变化。numeric_quality_summary.csv 则将这些逐样本指标汇总到模态级，用于判断异常是否是个别样本现象还是系统性问题。", lead="CSV 核验结构。")
    insert_text_before(p462, "本次三种模态的 finite_value_pct 均为 100%，mask_invalid_sample_count 均为 0，mask=1 但特征整行全零的情况均为 0。文本全零行较多是因为短文本在 50 个位置上产生空桶；视觉全零行主要来自无脸桶和 10 条整段失败样本，均由 mask=0 与状态字段共同解释。", lead="数值结论。")

    p463 = find_paragraph(doc, "4.6.3 官方接口交叉核验")
    insert_text_before(p463, "官方接口核验的范围被限定为字段、形状、sample_id 和标签映射。它回答的是“本队生成的数据能否以官方 aligned-50 接口被读取，以及样本和标签口径是否一致”，不回答“本队从原始视频提取的特征是否逐元素复现官方数值”。因此本节不报告特征向量之间的误差，也不把接口一致性写成识别准确率。", lead="核验范围。")

    p47 = find_paragraph(doc, "4.7 本问小结")
    insert_text_before(p47, "从数据文件角度看，问题一的最小可交付单元不是一个孤立的三维数组，而是由 sample_id、特征张量、mask、桶边界、源计数、标签和状态字段组成的完整记录。CSV 负责提供逐样本和逐桶的可读审计视图，NPZ 负责高效保存模型输入，schema、配置和运行日志负责固定解释规则。", lead="结果组织。")
    insert_text_before(p47, "这些结果为后续问题提供了两个直接接口：一方面，aligned-50 的三组特征可以按样本批量读取；另一方面，原始 mask 可以在后续缺失模态实验中作为观测基线。若模拟人为缺失，必须另建 synthetic_missing_mask，并在训练、验证和测试记录中区分原始缺失与人为缺失。", lead="后续衔接。")

    # Re-run-safe marker and basic metadata.
    doc.core_properties.title = "问题一求解及分析"
    doc.core_properties.subject = "复杂场景下多模态情感识别的特征建模与时间组织"
    doc.core_properties.comments = "扩展正文、CSV 字段解释和真实统计结果。"
    doc.save(str(OUTPUT))
    print(OUTPUT)


if __name__ == "__main__":
    main()
